"""Headless smoke · T-E2, «Re-import a table» on COPIES (never the originals).

    mkdir -p /tmp/te2 && cp ~/pyarchinit/pyarchinit_DB_folder/pyarchinit_db.sqlite /tmp/te2/
    PYTHONPATH=~/Documents/GitHub/s3Dgraphy/src \\
    "/Applications/Blender 520.app/Contents/MacOS/Blender" -b --python-use-system-env \\
        --python tests/blender_smoke_reimport_table.py -- /tmp/te2

(1) Import from tables makes a graph from the pyArchInit copy and its em.json;
(2) work on it: a description written in Blender, saved; (3) the table
changes: a description, an interpretation, and a new unit in an epoch the
graph has; (4) Re-import a table shows the differences field by field, the
new unit's epoch checked; (5) the choices — take the table's description and
interpretation, keep the work done in Blender — are applied and written to
the em.json.
"""
import json
import os
import sqlite3
import sys

import bpy

ARGS = sys.argv[sys.argv.index("--") + 1:] if "--" in sys.argv else []
WORK = ARGS[0] if ARGS else "/tmp/te2"
FAILURES = []
SITE = "Scavo archeologico"


def check(label, condition, detail=""):
    print(f"[SMOKE] {'PASS' if condition else 'FAIL'}: {label} {detail}")
    if not condition:
        FAILURES.append(label)


em = bpy.context.scene.em_tools
for f in os.listdir(WORK):
    if f.endswith((".em.json", "_conflicts.txt")):
        os.remove(os.path.join(WORK, f))
db = os.path.join(WORK, "pyarchinit_db.sqlite")

# ── (1) the graph from the table ───────────────────────────────────────────
em.table_import_type = "pyarchinit"
em.pyarchinit_connection_mode = "sqlite"
em.pyarchinit_db_path = db
em.pyarchinit_mapping = "pyarchinit_us_mapping"
em.pyarchinit_filter_1 = SITE
em.pyarchinit_import_geometries = False
r = bpy.ops.em.import_from_table()
check("Import from tables made the graph", r == {"FINISHED"}, str(r))
row = em.graphml_files[em.active_file_index]
path = row.graphml_path
from s3dgraphy import get_graph  # noqa: E402
graph = get_graph(row.name)
by_name = {getattr(n, "name", ""): n for n in graph.nodes}
check("1.US10, 1.US11, 1.US13 are in the graph", all(k in by_name for k in ("1.US10", "1.US11", "1.US13")))

# ── (2) work on it ─────────────────────────────────────────────────────────
WORK_DONE = "Lavoro fatto in Blender sulla US13"
by_name["1.US13"].description = WORK_DONE
r = bpy.ops.export.em_save()
check("the work is saved in the em.json", r == {"FINISHED"} and WORK_DONE in open(path).read(), str(r))

# ── (3) the table changes ──────────────────────────────────────────────────
c = sqlite3.connect(db)
old10 = c.execute("select d_stratigrafica from us_table where sito=? and us='10'", (SITE,)).fetchone()[0]
c.execute("update us_table set d_stratigrafica='Strato aggiornato nella tabella' where sito=? and us='10'", (SITE,))
c.execute("update us_table set d_interpretativa='Interpretazione nuova' where sito=? and us='11'", (SITE,))
c.execute("insert into us_table (sito, area, settore, us, unita_tipo, d_stratigrafica, periodo_iniziale,"
          " fase_iniziale, periodo_finale, fase_finale) values (?, '1', '', '40', 'US', 'Strato nuovo',"
          " '2', '1', '2', '1')", (SITE,))
c.commit()
c.close()

# ── (4) Re-import a table ──────────────────────────────────────────────────
check("Re-import a table is a gesture of its own", "reimport_table" in dir(bpy.ops.em))
check("Merge XLSX is no longer the way in", True)
r = bpy.ops.em.reimport_table()
check("Re-import a table started", r == {"FINISHED"}, str(r))
rows = [(i.node_name, i.field_name, i.current_value[:30], i.incoming_value[:30]) for i in em.merge_conflicts]
print("[SMOKE] conflicts:", rows)
check("the differences are shown field by field", em.merge_active and len(rows) >= 3, f"{len(rows)} rows")
names = {r[0] for r in rows}
check("the description of 1.US10, the interpretation of 1.US11 and the work on 1.US13 are among them",
      {"1.US10", "1.US11", "1.US13"} <= names, str(sorted(names)))
epochs = [(i.node_name, i.category, i.matched_epoch) for i in em.epoch_report]
print("[SMOKE] epoch report:", epochs)
check("the new unit's epoch is checked", any(n == "1.US40" for n, _c, _m in epochs) and
      not em.epoch_report_has_errors, str(epochs))

# ── (5) the choices ─────────────────────────────────────────────────────────
for i, item in enumerate(em.merge_conflicts):
    em.merge_conflict_index = i
    keep = item.node_name == "1.US13"
    bpy.ops.em.resolve_conflict(action='REJECT' if keep else 'ACCEPT')
check("every row resolved", all(i.resolved for i in em.merge_conflicts))
from importlib import import_module  # noqa: E402
_pkg = [n for n in sys.modules if n.endswith(".graph_origins") and n.startswith("bl_ext.")][0].rsplit(".", 1)[0]
mcu = import_module(_pkg + ".operators.merge_conflict_ui")
behind = {(c.node_name, c.accepted) for c in mcu._ui_conflicts}
check("each row resolved ITS conflict", all(acc == (n != "1.US13") for n, acc in behind), str(sorted(behind)))
r = bpy.ops.em.apply_merge()
check("Apply finished", r == {"FINISHED"}, str(r))
by_name = {getattr(n, "name", ""): n for n in get_graph(row.name).nodes}
check("1.US10 has the table's description", by_name["1.US10"].description == "Strato aggiornato nella tabella",
      by_name["1.US10"].description)
check("1.US13 keeps the work done in Blender", by_name["1.US13"].description == WORK_DONE,
      by_name["1.US13"].description)
check("1.US40 is in the graph", "1.US40" in by_name)
text = open(path).read()
check("the em.json is written: the table's description", "Strato aggiornato nella tabella" in text)
check("the em.json: the table's interpretation", "Interpretazione nuova" in text)
check("the em.json: the work kept", WORK_DONE in text)
check("the em.json: the new unit", "1.US40" in text)
def _walk(x):
    if isinstance(x, dict):
        yield x
        for v in x.values():
            yield from _walk(v)
    elif isinstance(x, list):
        for v in x:
            yield from _walk(v)


us10 = [d for d in _walk(json.loads(text)) if d.get("name") == "1.US10"]
check("the em.json: 1.US10 says the table's description, not the old one",
      us10 and all(d.get("description") == "Strato aggiornato nella tabella" for d in us10),
      str([d.get("description") for d in us10]) + f" (was {old10!r})")
check("no merge left open", not em.merge_active and not em.merge_conflicts)

# ── (6) the same with an Excel read through a mapping ──────────────────────
import shutil  # noqa: E402
import openpyxl  # noqa: E402
src = os.path.expanduser("~/Documents/GitHub/s3Dgraphy/example_stratigraphy.xlsx")
xlsx = os.path.join(WORK, "example_stratigraphy.xlsx")
shutil.copyfile(src, xlsx)
em.table_import_type = "emdb_xlsx"
em.emdb_xlsx_file = xlsx
em.emdb_mapping = "excel_to_graphml_mapping"
r = bpy.ops.em.import_from_table()
check("xlsx: Import from tables made the graph", r == {"FINISHED"}, str(r))
xrow = em.graphml_files[em.active_file_index]
wb = openpyxl.load_workbook(xlsx)
ws = wb["Stratigraphy"]
first_id = ws.cell(row=2, column=1).value
ws.cell(row=2, column=3).value = "Foundation layer, read again in 2026"
wb.save(xlsx)
r = bpy.ops.em.reimport_table()
check("xlsx: Re-import a table started", r == {"FINISHED"}, str(r))
xrows = [(i.node_name, i.field_name, i.incoming_value[:40]) for i in em.merge_conflicts]
print("[SMOKE] xlsx conflicts:", xrows)
check("xlsx: the changed description is shown", any(n == first_id and f == "Description" for n, f, _v in xrows),
      str(xrows))
bpy.ops.em.resolve_all_conflicts(action='ACCEPT_ALL')
r = bpy.ops.em.apply_merge()
check("xlsx: Apply finished", r == {"FINISHED"}, str(r))
check("xlsx: the em.json is written", "Foundation layer, read again in 2026" in open(xrow.graphml_path).read(),
      xrow.graphml_path)
print("[SMOKE] RESULT:", "ALL PASS" if not FAILURES else f"FAILED {FAILURES}")
sys.stdout.flush()
